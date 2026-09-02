from __future__ import annotations

import io
import json

import pytest
from docx import Document

from cairn.server import report_template_service


def _template_bytes() -> bytes:
    document = Document()
    summary = document.add_table(rows=2, cols=4)
    for cell, value in zip(summary.rows[0].cells, ("序号", "漏洞名称", "风险级别", "是否修复")):
        cell.text = value
    for cell, value in zip(summary.rows[1].cells, ("01", "示例", "中风险", "未修复")):
        cell.text = value

    document.add_paragraph("附件：渗透检测结果详情")
    detail = document.add_table(rows=9, cols=4)
    labels = (
        "01",
        "漏洞名称",
        "漏洞描述",
        "修复情况",
        "漏洞位置/URL",
        "漏洞证明",
        "修复建议",
        "修复情况说明",
        "修复验证说明",
    )
    for row, label in zip(detail.rows, labels):
        row.cells[0].text = label
    detail.rows[1].cells[2].text = "风险程度"

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _nested_plan() -> dict:
    return {
        "summary_table_index": 0,
        "summary_header_row": 0,
        "summary_data_start_row": 1,
        "summary_columns": {"number": 0, "title": 1, "severity": 2, "status": 3},
        "detail_table_indices": [1],
        "detail_value_column": {"1": 1},
        "detail_field_rows": {
            "1": {
                "number": 0,
                "title": 1,
                "severity": 1,
                "description": 2,
                "status": 3,
                "location": 4,
                "proof": 5,
                "remediation": 6,
                "fix_note": 7,
                "retest_note": 8,
            }
        },
        "detail_field_columns": {
            "1": {
                "number": 0,
                "title": 1,
                "severity": 3,
                "description": 1,
                "status": 1,
                "location": 1,
                "proof": 1,
                "remediation": 1,
                "fix_note": 1,
                "retest_note": 1,
            }
        },
    }


def _vulnerabilities(count: int = 5) -> list[dict]:
    return [
        {
            "title": f"漏洞 {index}",
            "severity": "high",
            "status": "confirmed",
            "description": f"漏洞描述 {index}",
            "evidence": [f"已确认事实 {index}"],
            "proof_packets": [],
            "screenshots": [],
        }
        for index in range(1, count + 1)
    ]


def _model_rows(count: int = 5) -> str:
    rows = []
    for index in range(1, count + 1):
        rows.append(
            {
                "fields": {
                    "number": str(index),
                    "title": f"漏洞 {index}",
                    "severity": "高风险",
                    "description": f"漏洞描述 {index}",
                    "status": "已确认",
                    "location": f"/target/{index}",
                    "proof": f"模型整理的事实证明 {index}",
                    "remediation": [f"修复建议 {index}"],
                    "fix_note": "待责任单位填写",
                    "retest_note": "待复测",
                }
            }
        )
    return json.dumps({"vulnerabilities": rows}, ensure_ascii=False)


def test_compose_supports_nested_agent_plan_and_fills_all_detail_tables(monkeypatch):
    monkeypatch.setattr(report_template_service, "_resolve_report_composer_profile", lambda: object())
    monkeypatch.setattr(report_template_service, "_request_model_json", lambda **_kwargs: _model_rows())

    body = report_template_service.compose_template_report(
        _template_bytes(),
        {"plan": _nested_plan()},
        {"id": "proj_test", "name": "测试项目"},
        _vulnerabilities(),
    )

    document = Document(io.BytesIO(body))
    assert len(document.tables) == 6
    summary = document.tables[0]
    assert len(summary.rows) == 6
    assert [cell.text for cell in summary.rows[1].cells] == ["01", "漏洞 1", "高风险", "未修复"]
    assert [cell.text for cell in summary.rows[5].cells] == ["05", "漏洞 5", "高风险", "未修复"]

    for index, table in enumerate(document.tables[1:], start=1):
        assert table.rows[0].cells[0].text == f"{index:02d}"
        assert table.rows[1].cells[1].text == f"漏洞 {index}"
        assert table.rows[1].cells[3].text == "高风险"
        assert table.rows[2].cells[1].text == f"漏洞描述 {index}"
        assert table.rows[3].cells[1].text == "未修复"
        assert table.rows[4].cells[1].text == f"/target/{index}"
        assert f"模型整理的事实证明 {index}" in table.rows[5].cells[1].text
        assert "当前事实未保存原始请求/响应数据包或证明截图" in table.rows[5].cells[1].text
        assert table.rows[6].cells[1].text == f"修复建议 {index}"


def test_compose_embeds_real_packets_without_claiming_they_are_missing(monkeypatch):
    monkeypatch.setattr(report_template_service, "_resolve_report_composer_profile", lambda: object())
    monkeypatch.setattr(report_template_service, "_request_model_json", lambda **_kwargs: _model_rows(1))
    vulnerabilities = _vulnerabilities(1)
    vulnerabilities[0]["proof_packets"] = [
        {
            "title": "目录越界验证",
            "request": "GET /download?path=../../etc/passwd HTTP/1.1",
            "response": "HTTP/1.1 200 OK\nroot:x:0:0",
            "note": "响应命中系统账号文件",
        }
    ]

    body = report_template_service.compose_template_report(
        _template_bytes(), {"plan": _nested_plan()}, {"name": "测试项目"}, vulnerabilities
    )

    proof = Document(io.BytesIO(body)).tables[1].rows[5].cells[1].text
    assert "GET /download?path=../../etc/passwd HTTP/1.1" in proof
    assert "HTTP/1.1 200 OK" in proof
    assert "当前事实未保存" not in proof


def test_compose_rejects_missing_model_rows_instead_of_silently_omitting_findings(monkeypatch):
    monkeypatch.setattr(report_template_service, "_resolve_report_composer_profile", lambda: object())
    monkeypatch.setattr(report_template_service, "_request_model_json", lambda **_kwargs: _model_rows(1))

    with pytest.raises(RuntimeError, match="项目当前有 2 条"):
        report_template_service.compose_template_report(
            _template_bytes(), {"plan": _nested_plan()}, {"name": "测试项目"}, _vulnerabilities(2)
        )


def test_inspect_template_normalizes_single_table_nested_plan(monkeypatch):
    monkeypatch.setattr(report_template_service, "_resolve_report_composer_profile", lambda: object())
    monkeypatch.setattr(
        report_template_service,
        "_request_model_json",
        lambda **_kwargs: json.dumps(_nested_plan(), ensure_ascii=False),
    )

    analysis = report_template_service.inspect_template(_template_bytes())

    assert analysis["plan"]["detail_value_column"] == 1
    assert analysis["plan"]["detail_field_rows"]["proof"] == 5
    assert analysis["plan"]["detail_field_columns"]["severity"] == 3
